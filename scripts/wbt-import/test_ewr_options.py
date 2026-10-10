"""The optional [EWR options] sheet (engine >= 1.77.0, issue #455, docs/model.md 2.9f).

    .venv/bin/python -m unittest discover -s scripts/wbt-import

read_ewr_options() turns the sheet's defined names into settings.ewrDailySource; a workbook
without them imports as before. The browser importer (frontend/src/lib/spreadsheet/import/ewrOptions.ts)
must give the same values and notes: its parity test reads the committed variant fixture
(fixtures/synthetic_b023.ewr-options.*), which this test keeps in step with the generator.
"""

import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))

import make_synthetic_workbook as gen  # noqa: E402
from extract_project import Workbook, extract, read_ewr_options  # noqa: E402

POINTS = [0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9, 0.99]


def sheet(method="TAB file", scaling="MAR ratio", mar=100.0, area=None, tab=None, natural=None, reserve=None, points=None) -> Workbook:
    """An in-memory workbook holding only an [EWR options] sheet with the given values (None = blank)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EWR options"
    ws["A1"], ws["A2"], ws["A3"], ws["A4"] = method, scaling, mar, area
    for i, v in enumerate(tab if tab is not None else [0.5] * 12):
        ws.cell(row=10 + i, column=1, value=v)
    for j, p in enumerate(points if points is not None else POINTS):
        ws.cell(row=5, column=3 + j, value=p)
    for name, grid, top in (("zEwrOpt_NaturalPct", natural, 10), ("zEwrOpt_ReservePct", reserve, 30)):
        for i, row in enumerate(grid if grid is not None else [[1.0] * 10] * 12):
            for j, v in enumerate(row):
                ws.cell(row=top + i, column=3 + j, value=v)
    names = {
        "zEwrOpt_Method": "$A$1",
        "zEwrOpt_Scaling": "$A$2",
        "zEwrOpt_TableMar": "$A$3",
        "zEwrOpt_TableArea": "$A$4",
        "zEwrOpt_TabM3s": "$A$10:$A$21",
        "zEwrOpt_PctPoints": "$C$5:$L$5",
        "zEwrOpt_NaturalPct": "$C$10:$L$21",
        "zEwrOpt_ReservePct": "$C$30:$L$41",
    }
    for n, r in names.items():
        wb.defined_names[n] = DefinedName(n, attr_text=f"'EWR options'!{r}")
    return Workbook(wb)


class ReadEwrOptions(unittest.TestCase):
    def test_no_sheet_no_setting(self):
        wb = Workbook(openpyxl.Workbook())
        self.assertIsNone(read_ewr_options(wb))
        self.assertEqual(wb.notes, [])

    def test_a_tab_source(self):
        wb = sheet(tab=[0.1 * (i + 1) for i in range(12)])
        out = read_ewr_options(wb)
        self.assertEqual(out["method"], "tab")
        self.assertEqual(out["scaling"], "mar")
        self.assertEqual(out["tableMarMm3"], 100.0)
        self.assertIsNone(out["tableAreaKm2"])
        self.assertAlmostEqual(out["tabM3s"][11], 1.2)
        self.assertEqual(wb.notes, ["[EWR options] the daily EWR at the outlet: the DRM TAB file, scaled by the MAR ratio (docs/model.md 2.9f)"])

    def test_names_are_read_case_insensitively_and_blank_is_the_default(self):
        out = read_ewr_options(sheet(method="percentile TABLES", scaling="Area Ratio", area=40))
        self.assertEqual((out["method"], out["scaling"], out["tableAreaKm2"]), ("percentile", "area", 40.0))
        self.assertEqual(read_ewr_options(sheet(method=None, scaling=None))["method"], "pragmatic")
        # A blank scaling is the area ratio (issue #90 B2), so a TAB file needs the table area.
        wb = sheet(method="TAB file", scaling=None, mar=None, area=40, tab=[0.5] * 12)
        out = read_ewr_options(wb)
        self.assertEqual((out["method"], out["scaling"], out["tableAreaKm2"]), ("tab", "area", 40.0))
        self.assertEqual(wb.notes, ["[EWR options] the daily EWR at the outlet: the DRM TAB file, scaled by the area ratio (docs/model.md 2.9f)"])

    def test_an_unknown_method_or_scaling_warns_and_falls_back(self):
        wb = sheet(method="Monthly", scaling="Volume")
        out = read_ewr_options(wb)
        self.assertEqual((out["method"], out["scaling"]), ("pragmatic", "area"))
        self.assertTrue(wb.notes[0].startswith('WARNING: [EWR options] the EWR method "Monthly"'))
        self.assertEqual(wb.notes[1], 'WARNING: [EWR options] the scaling "Volume" is not MAR ratio or Area ratio: imported as Area ratio')

    def test_a_method_missing_what_it_needs_imports_as_pragmatic_and_keeps_the_values(self):
        wb = sheet(method="Percentile tables", scaling="Area ratio", natural=[[1.0] * 9 + [None]] + [[1.0] * 10] * 11)
        out = read_ewr_options(wb)
        self.assertEqual(out["method"], "pragmatic")
        self.assertIsNone(out["naturalPctM3s"])
        self.assertIsNotNone(out["reservePctM3s"])
        self.assertEqual(
            wb.notes[0],
            "WARNING: [EWR options] the DRM percentile tables needs the natural flow percentile table, the table area "
            "(a number for every cell, none below 0): imported as the pragmatic EWR",
        )

    def test_a_negative_or_text_value_drops_that_table(self):
        self.assertIsNone(read_ewr_options(sheet(tab=[0.5] * 11 + [-1]))["tabM3s"])
        self.assertIsNone(read_ewr_options(sheet(tab=[0.5] * 11 + ["x"]))["tabM3s"])
        self.assertIsNone(read_ewr_options(sheet(mar=0))["tableMarMm3"])

    def test_other_percentile_points_warn(self):
        wb = sheet(points=[10, 20, 30, 40, 50, 60, 70, 80, 90, 99])
        read_ewr_options(wb)
        self.assertTrue(wb.notes[0].startswith("WARNING: [EWR options] the percentile points are not 0.1"))


class Fixture(unittest.TestCase):
    def test_the_committed_variant_matches_a_regenerated_one(self):
        with tempfile.TemporaryDirectory() as tmp:
            fresh = Path(tmp) / gen.EWR_OPTIONS_WORKBOOK
            gen.write_workbook(fresh, with_ewr_options=True)
            committed = (gen.FIXTURES / f"{gen.EWR_OPTIONS_STEM}.json").read_text(encoding="utf-8")
            self.assertEqual(gen.ewr_options_output(fresh), committed)
            self.assertEqual(gen.ewr_options_output(gen.FIXTURES / gen.EWR_OPTIONS_WORKBOOK), committed)

    def test_the_sheet_only_adds_the_setting(self):
        base, _, base_notes = extract(gen.FIXTURES / gen.WORKBOOK_NAME)
        variant, _, notes = extract(gen.FIXTURES / gen.EWR_OPTIONS_WORKBOOK)
        self.assertNotIn("ewrDailySource", base["settings"])
        self.assertEqual(variant["settings"]["ewrDailySource"]["method"], "percentile")
        rest = {k: v for k, v in variant["settings"].items() if k != "ewrDailySource"}
        self.assertEqual(json.dumps(rest, sort_keys=True), json.dumps(base["settings"], sort_keys=True))
        self.assertEqual([n for n in notes if "[EWR options]" not in n], base_notes)


if __name__ == "__main__":
    unittest.main()
