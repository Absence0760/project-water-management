"""Generate the synthetic b023 workbook fixture and the Python importer's output for it.

    python scripts/wbt-import/make_synthetic_workbook.py [OUTDIR]

Writes, into OUTDIR (default scripts/wbt-import/fixtures/):

  synthetic_b023.xlsx                        a b023-layout workbook of an INVENTED catchment
  synthetic_b023.project.json                extract_project.py's project for it
  synthetic_b023.notes.txt                   extract_project.py's notes, one per line
  synthetic_b023.gauge-reference.project.json   the same with --gauge-as-reference
  synthetic_b023.gauge-reference.notes.txt      --gauge-scaling-from 2021-10-01 --gauge-scale-factor 0.8
  synthetic_b023.run-of-river.project.json      the same with --run-of-river (the browser importer's
  synthetic_b023.run-of-river.notes.txt         run-of-river option checks its parity against these)

Everything in the workbook is made up: the farm, crop and gauge names, the
areas, dams, crop factors and every daily value (a seeded random generator).
Nothing is copied from a client workbook; only the b023 layout (sheet names,
named ranges, which cell holds what) follows the tool. It is the fixture for
the in-browser importer's parity test (WP-1.31): the TypeScript port must turn
this workbook into the same project and notes as the Python importer.

The workbook carries plain values in every cell the importer reads, as if
Excel had calculated and saved it: openpyxl writes a formula without a cached
value, and the importer opens the file with data_only=True, so a formula
there would read as blank. The formula-shaped parts of b023 that no importer
evaluates are still there as formulas: the [Transfers] formula row (row 7),
with its text copy in row 8 (zTransfers_FormulasAsTxt) as plain strings.

The output is deterministic: a fixed seed, fixed dates and fixed zip
timestamps, so regenerating gives the same importer output (test_synthetic.py
checks it) and, with the same openpyxl version, the same workbook bytes.
See scripts/wbt-import/README.md, "The synthetic workbook", for which part of
the fixture exercises which importer branch.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import math
import random
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from extract_project import extract  # noqa: E402

FIXTURES = HERE / "fixtures"
STEM = "synthetic_b023"
WORKBOOK_NAME = f"{STEM}.xlsx"
GAUGE_REFERENCE_ARGS = {"gauge_reference": True, "scaling_from": "2021-10-01", "scale_factor": 0.8}

SEED = 20260925
START = dt.date(2019, 10, 1)
END = dt.date(2022, 9, 30)  # 1096 days: three hydrological years
EXTRA_DAYS_AFTER_END = 2  # dated rows after zFlowData_DateE_DateSeries, which the importer must not read
FIXED_TIME = dt.datetime(2026, 1, 1)
ZIP_TIME = (2026, 1, 1, 0, 0, 0)

MONTHS_WY = ["Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"]
MONTH_NUMS_WY = [10, 11, 12, 1, 2, 3, 4, 5, 6, 7, 8, 9]
DAYS_WY = [31, 30, 31, 31, 28.25, 31, 30, 31, 30, 31, 31, 30]

# --------------------------------------------------------------------------
# The invented catchment. Names are NATO-alphabet placeholders so they can't
# resemble a real farm or place.

# [Network], upstream first: (element name as typed, type, upstream elements).
NETWORK: list[tuple[str, str, list[str]]] = [
    ("Alpha Farm", "Farm", []),
    ("Bravo Farm", "Farm", ["Alpha Farm"]),
    ("Charlie Farm", "Farm", []),
    ("Delta Farm", "Farm", []),  # not in [Farm spec]
    ("Echo  Farm", "Farm", ["Bravo Farm", "Charlie Farm", "Delta Farm"]),  # double space: names are whitespace-normalised
    ("Midway Gauge", "Gauge", ["Echo Farm"]),  # a gauge inside the network
    ("Foxtrot Farm", "Farm", ["Midway Gauge"]),
    ("Golf\nFarm", "Farm", []),  # a line break typed in the cell (Alt+Enter): read as one space (issue #385)
    ("India Farm", "Farm", ["Golf Farm"]),  # a "dummy dam": pumps from the river
    ("Hotel Farm", "Farm", ["Foxtrot Farm", "India Farm"]),  # no Element sheet
    ("Outlet Gauge", "Gauge", ["Hotel Farm"]),
]
OUTFLOW_GAUGE = "Outlet Gauge"
NO_ELEMENT_SHEET = {"Hotel Farm"}

# [Farm spec]: hi/lo MAP area (km2), total area (km2), external fragmentation,
# % upstream inflow above dam (as b023 holds it: its formula sends that share
# past the dam, so the importer stores 1 − it, docs/model.md §3 Q1),
# % farm runoff into dam, dam capacity (m3),
# initial %, min % for transfers, diversion back to dam (m3/day), irrigation return flow %.
FARM_SPEC: list[dict[str, Any]] = [
    {"name": "Alpha Farm", "hi": 6.0, "lo": 4.0, "total": 10.0, "ext": 0.12, "up": 1, "runoff": 0.3, "cap": 250000, "init": 0.5, "min": 0, "divert": 8640, "ret": 0.2},
    {"name": "Bravo Farm", "hi": 8.5, "lo": 5.5, "total": 14.0, "ext": 0.17, "up": 0.5, "runoff": 0.25, "cap": 600000, "init": 0.6, "min": 0.3, "divert": 21600, "ret": 0.2},
    # A "Specific" method farm with no external fragmentation and no dam: blanks read as 0.
    {"name": "Charlie Farm", "hi": 3.2, "lo": 6.8, "total": 10.0, "ext": None, "up": 1, "runoff": 0, "cap": None, "init": None, "min": None, "divert": 0, "ret": 0},
    {"name": "Echo Farm", "hi": 9.0, "lo": 7.0, "total": 16.0, "ext": 0.2, "up": 1, "runoff": 0.4, "cap": 1500000, "init": 0.5, "min": 0.25, "divert": 21600, "ret": 0.15},
    # The diversion typed as text, as a pasted cell can be.
    {"name": "Foxtrot Farm", "hi": 5.0, "lo": 5.0, "total": 10.0, "ext": 0.13, "up": 1, "runoff": 0.1, "cap": 80000, "init": 0.4, "min": 0, "divert": "4320", "ret": 0},
    # Total area 8.5 km2 is not hi + lo (8.0): the importer keeps both as typed.
    {"name": "Golf Farm", "hi": 4.4, "lo": 3.6, "total": 8.5, "ext": 0.11, "up": 1, "runoff": 0.9, "cap": 120000, "init": 1.0, "min": 0.2, "divert": 0, "ret": 0.3},
    {"name": "Hotel Farm", "hi": 7.0, "lo": 8.0, "total": 15.0, "ext": 0.16, "up": 0.75, "runoff": 0.2, "cap": 45000, "init": 0.5, "min": 0, "divert": 8640, "ret": 0.1},
    # b023's stand-in for a unit that pumps straight from the river: a 20 m3
    # pool that takes all the upstream inflow (0 % in b023's terms), with a large diversion capacity.
    # The importer warns "probable run-of-river" (issue #54, 2d).
    {"name": "India Farm", "hi": 1.5, "lo": 1.0, "total": 2.5, "ext": 0.04, "up": 0, "runoff": 0.2, "cap": 20, "init": 1.0, "min": 0, "divert": 12960, "ret": 0.1},
]
FRAGMENTATION_METHOD = "Specific"
HI_LO_SPLIT = (0.8, 0.2)

APAN_MM = [190.0, 240.0, 285.0, 300.0, 250.0, 225.0, 150.0, 95.0, 65.0, 70.0, 95.0, 135.0]  # Oct..Sep
EFFECTIVE_RAIN = 0.7
CROPS: list[tuple[str, list[float]]] = [
    ("Orchard A", [0.45, 0.5, 0.55, 0.6, 0.6, 0.55, 0.5, 0.45, 0.4, 0.4, 0.4, 0.42]),
    ("Vegetables B", [0.7, 0.85, 1.0, 0.3, 0.5, 0.75, 0.9, 1.0, 0, 0, 0.35, 0.5]),
    ("Pasture C", [0.7, 0.75, 0.8, 0.8, 0.8, 0.8, 0.7, 0.6, 0.5, 0.5, 0.5, 0.6]),
    ("Vines D", [0.25, 0.35, 0.45, 0.5, 0.45, 0.3, 0.15, 0, 0, 0, 0.1, 0.2]),
    # A crop no farm grows, with a lone 0 (Dec) and a spike above 1 (Mar): the crop-table check warns (issue #289).
    ("Fodder E", [0.5, 0.5, 0, 0.5, 0.5, 1.2, 0.5, 0.5, 0.5, 0.5, 0.5, 0.5]),
    ("Pasture F", [0.7, 0.75, 0.8, 0.8, 0.8, 0.8, 0.7, 0.6, 0.5, 0.5, 0.5, 0.6]),  # Pasture C's row pasted: warned as copied
]
# [Crop demand] names as typed, where they differ from CROPS: a stray C1 control character
# (U+009F, as a paste from another program can leave) the importers read as a space (issue #385).
CROP_AS_TYPED = {"Vines D": "Vines\x9fD"}
# [Farm demand] crop columns: the six crops plus one that [Crop demand] doesn't define.
FARM_DEMAND_CROPS = [c[0] for c in CROPS] + ["Hops X"]
# Farm (as typed in [Farm demand]) -> area m2 per FARM_DEMAND_CROPS column. None = a blank row.
FARM_DEMAND: list[tuple[str, list[float | None]] | None] = [
    ("Alpha Farm", [120000, 0, 80000, 0, 0, 0, 0]),
    ("Bravo Farm", [0, 150000, None, 60000, 0, 0, 0]),
    ("Charlie Farm", [0, 0, 0, 0, 0, 0, 0]),
    None,
    ("Echo Farm ", [200000, 50000, 0, 0, 0, 0, 25000]),  # trailing space; grows the undefined crop
    ("Foxtrot Farm", [0, 0, 90000, 0, 0, 0, 0]),
    (" Golf Farm", [30000, 0, 0, 45000, 0, 0, 0]),
    ("Hotel Farm", [0, 70000, 0, 0, 0, 0, 0]),
    ("India Farm", [0, 0, 0, 40000, 0, 0, 0]),
    ("Kilo Farm", [50000, 0, 0, 0, 0, 0, 0]),  # not in [Network]
]
# [Farm demand] gross demand typed over the formula (m³/day, every month), as a
# workbook does for a town's potable demand: the importer warns (issue #54).
TYPED_GROSS_M3_DAY = {"Charlie Farm": 150.0}

# [Transfers] "Draw from dam" columns: (from, transfer to, months cell, max m3/s, min %, capacity m3/day).
TRANSFERS: list[tuple[str, str, Any, float, float, float]] = [
    ("Bravo Farm", "Alpha Farm", "1,2,3,11,12", 0.25, 0.25, 21600),
    ("Foxtrot Farm", "Echo Farm", "11,12", 0.1, 0, 8640),  # substring match adds Jan, Feb (M1)
    ("Hotel Farm", "Golf Farm", 7, 0.05, 0.25, 4320),  # a single month, stored as a number
    ("Bravo Farm", "Charlie Farm", " 5, 6 ,7", 0.15, 0.3, 12960),  # a second rule on one dam: priority 1 after 0
    ("Echo Farm", "--", "1,2,3", 0.1, 0.25, 8640),  # no destination: not a transfer
    ("Golf Farm", "Foxtrot Farm", "1,2", 0, 0.2, 0),  # max rate 0: not a transfer
    ("Alpha Farm", "Zulu Farm", "6,7,8", 0.05, 0, 4320),  # an element not in [Network]
    ("Foxtrot Farm", "Golf Farm", "10", 0.05, 0, 4320),  # its draw formula is =0: switched off (issue #54)
    # From India's dummy dam into Delta, which has no dam and no demand: a canal off-take (engine 1.14.0), switched off.
    ("India Farm", "Delta Farm", "1,2,3,4,5,6,7,8,9,10,11,12", 0.02, 0, 1728),
]
# Transfers (by index) whose draw formula is the constant 0: the workbook never moves that water.
TRANSFERS_OFF = {7, 8}
# InOut columns (one per farm): the hand-written formulas that move the drawn
# water. {Name} is replaced by that transfer's From column letter. All follow
# "add to the destination, subtract from the source" except Echo's, which
# delivers 90 % of the Foxtrot draw (a conveyance loss no transfer rule has).
IN_OUT: list[tuple[str, str]] = [
    ("Alpha Farm", "={T0}7-{T6}7"),
    ("Bravo Farm", "=-{T0}7-{T3}7"),
    ("Charlie Farm", "={T3}7"),
    ("Echo Farm", "={T1}7*0.9-{T4}7"),
    ("Foxtrot Farm", "={T5}7-{T1}7-{T7}7"),
    ("Golf Farm", "={T2}7+{T7}7-{T5}7"),
    ("Hotel Farm", "=-{T2}7"),
    ("Delta Farm", "={T8}7"),
    ("India Farm", "=-{T8}7"),
]

EWR_M3_DAY = [5200, 3100, 900, 700, 1300, 1700, 2400, 8800, 14500, 19800, 21000, 12600]  # Oct..Sep
USE_FLOW = 2  # [Flow data] rUseFlow: 1 Pitman, 2 gauge, 3 logger
CALIBRATION = {
    "D7": 10,
    "D8": 2019,
    "D9": 2021,
    "D11": dt.datetime(2019, 10, 1),  # zCalibration_Date1
    "D12": dt.datetime(2021, 9, 30),  # zCalibration_DateN
    "H8": 2,  # rain threshold (mm)
    "H9": 0.08,  # a
    "H10": 1.25,  # b
    "H11": "10,11,12,3,",  # summer months, trailing comma; substring match adds Jan, Feb (M1)
    "H12": 0.15,  # summer factor
    "H13": 0.9,  # winter factor
    "H16": 0.85,  # shift peak index lo
    "H17": 3.2,  # shift peak index hi
    "H18": 1.5,  # min ratio rain/base to reset base
    "H19": 8,  # winter starts day+1 if rain >=
    "H20": 30,  # winter starts day+0 if rain >=
    "H22": OUTFLOW_GAUGE,
    "AD4": 180000,  # base for peak flow
    "AD5": 1800000,  # amplitude (m3/day)
    "AD6": 40,  # max days in curve
    "AD7": 0.9,
}
BASE_FLOW_INITIAL = 320  # [Flow data] "reserved" row, Underlying Base Flow column

# Rain: a summer-rainfall catchment (wet Oct..Mar).
WET_MONTHS = {10, 11, 12, 1, 2, 3}
# Rain-quality cases (issue #2):
ZERO_RUN_WET = (dt.date(2021, 1, 5), dt.date(2021, 3, 20))  # logger outage recorded as zeros: flagged
ZERO_RUN_DRY = (dt.date(2020, 5, 1), dt.date(2020, 8, 31))  # a genuine winter dry spell: not flagged
RAIN_BLANK = (dt.date(2020, 2, 10), dt.date(2020, 2, 14))  # blank cells: None
RAIN_TEXT_DAY = dt.date(2020, 3, 3)  # a text cell ("err"): None
GAUGE_BLANK = (dt.date(2020, 11, 1), dt.date(2020, 11, 20))
GAUGE_TEXT_DAY = dt.date(2021, 2, 2)
GAUGE_SCALED_FROM = dt.date(2021, 10, 1)  # the gauge record is scaled by GAUGE_SCALE from here
GAUGE_SCALE = 0.8
LOGGER_FROM = dt.date(2021, 6, 1)
PITMAN_DAYS = 120


# --------------------------------------------------------------------------


def ref(c1: str, r1: int, c2: str | None = None, r2: int | None = None) -> str:
    a = f"${c1}${r1}"
    return a if c2 is None else f"{a}:${c2}${r2}"


def define(wb, name: str, sheet: str, rng: str) -> None:
    quoted = f"'{sheet}'" if re.search(r"[^A-Za-z0-9_]", sheet) else sheet
    wb.defined_names[name] = DefinedName(name, attr_text=f"{quoted}!{rng}")


def put(ws, cell: str, value: Any) -> None:
    ws[cell] = value


def put_text(ws, cell: str, value: str) -> None:
    """A string that starts with '=' stored as text, not as a formula."""
    c = ws[cell]
    c.value = value
    c.data_type = "s"


def col(c: str, offset: int) -> str:
    return get_column_letter(column_index_from_string(c) + offset)


def days() -> list[dt.date]:
    n = (END - START).days + 1 + EXTRA_DAYS_AFTER_END
    return [START + dt.timedelta(days=i) for i in range(n)]


def within(d: dt.date, span: tuple[dt.date, dt.date]) -> bool:
    return span[0] <= d <= span[1]


def farm_names() -> list[str]:
    return [re.sub(r"\s+", " ", n) for n, t, _ in NETWORK if t == "Farm"]


# --------------------------------------------------------------------------


def app_settings(wb) -> None:
    ws = wb.create_sheet("AppSettings")
    put(ws, "C8", "Water Balance Tool b023 (synthetic)")
    put(ws, "B95", "Month number")
    put(ws, "C95", "Month label")
    put(ws, "D95", "Month Days")
    labels = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    month_days = [31, 28.25, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    for i, (lab, d) in enumerate(zip(labels, month_days)):
        put(ws, f"B{96 + i}", i + 1)
        put(ws, f"C{96 + i}", lab)
        put(ws, f"D{96 + i}", d)
    put(ws, "B108", "-- do NOT delete this row !  Insert before")
    put(ws, "C108", "--")
    put(ws, "D108", "--")
    define(wb, "rAppSet_MonthLbls", "AppSettings", ref("C", 96, "C", 107))
    define(wb, "zAppSet_MonthDays", "AppSettings", ref("D", 96, "D", 107))


def home(wb) -> None:
    # [Home] calculation window: the whole flow record here, so the run covers it all (issue #54 model_window).
    ws = wb.create_sheet("Home")
    put(ws, "B1", "Home")
    put(ws, "D14", "Calc. from")
    put(ws, "D15", "Calc. to")
    put(ws, "E14", dt.datetime.combine(START, dt.time()))
    put(ws, "E15", dt.datetime.combine(END, dt.time()))
    define(wb, "zHome_CalcDate1", "Home", ref("E", 14))
    define(wb, "zHome_CalcDateN", "Home", ref("E", 15))


def templates(wb) -> None:
    for name, headers in (("FarmTemplate", FARM_HEADERS), ("GaugeTemplate", GAUGE_HEADERS)):
        ws = wb.create_sheet(name)
        put(ws, "B15", "Reserved")
        put(ws, "B16", "Reserved (init)")
        put(ws, "B17", "Formula row")
        put(ws, "E19", "Date")
        for c, h in headers.items():
            put(ws, f"{c}19", h)
        put(ws, "E20", "init")


def network(wb) -> None:
    ws = wb.create_sheet("Network")
    put(ws, "J20", "Synthetic catchment: invented element names and values (scripts/wbt-import/make_synthetic_workbook.py)")
    put(ws, "C26", "Gauge Names")
    put(ws, "D26", "TransferFarms")
    put(ws, "J26", "Element name")
    put(ws, "K26", "Element Type")
    put(ws, "L26", "Has transfer?")
    put(ws, "M26", "|")
    for i in range(7):
        put(ws, f"{col('N', i)}26", f"Upstream \nelement-{i + 1}")
    put(ws, "U26", "|")
    has_transfer = {t[0] for t in TRANSFERS} | {t[1] for t in TRANSFERS}
    r = 27
    for name, etype, ups in NETWORK:
        clean = re.sub(r"\s+", " ", name)
        if etype == "Gauge":
            put(ws, f"C{r}", clean)
        if clean in has_transfer:
            put(ws, f"D{r}", clean)
        put(ws, f"J{r}", name)
        put(ws, f"K{r}", etype)
        put(ws, f"L{r}", "Yes" if clean in has_transfer else "No")
        put(ws, f"M{r}", "|")
        for i, u in enumerate(ups):
            put(ws, f"{col('N', i)}{r}", u)
        put(ws, f"U{r}", "|")
        r += 1
    sentinel = r
    for c in ["C", "D", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U"]:
        put(ws, f"{c}{sentinel}", "--")
    put(ws, f"J{sentinel}", "-- do NOT delete this row !  Insert here")
    put(ws, f"J{sentinel + 2}", "Network Outflow Gauge:")
    put(ws, f"J{sentinel + 3}", OUTFLOW_GAUGE)
    define(wb, "zNetwork_ElementNameLst", "Network", ref("J", 26, "J", sentinel))
    define(wb, "zNetwork_ElementTypeLst", "Network", ref("K", 26, "K", sentinel))
    define(wb, "zNetwork_ElementGaugeNameLst", "Network", ref("C", 26, "C", sentinel))
    define(wb, "zNetwork_FarmsTransferLst", "Network", ref("D", 26, "D", sentinel))
    define(wb, "zNetwork_HasTransferYnLst", "Network", ref("L", 26, "L", sentinel))
    define(wb, "zNetwork_UpstreamTbl", "Network", ref("M", 26, "U", sentinel))
    define(wb, "zNetwork_OutflowGauge", "Network", ref("J", sentinel + 3))


def farm_spec(wb) -> None:
    s = "Farm spec"
    ws = wb.create_sheet(s)
    put(ws, "A1", "Farm spec")
    put(ws, "B16", "Fragmentation tolerance: ")
    put(ws, "C16", 0.0002)
    put(ws, "M3", "Method")
    for i, m in enumerate(["Area", "Hi/Lo", "Specific"]):
        put(ws, f"M{4 + i}", m)
    put(ws, "I26", "(Pitman) Hi MAP flow %")
    put(ws, "J26", "Lo MAP flow %")
    put(ws, "I27", HI_LO_SPLIT[0])
    put(ws, "J27", HI_LO_SPLIT[1])
    put(ws, "M26", "Selected method:")
    put(ws, "M27", FRAGMENTATION_METHOD)
    headers = {
        "D": "Farm name",
        "E": "Type check",
        "F": "Sequence check",
        "G": "Total Area\nFarms\n\n(km²)",
        "H": "Area based fragmen\ntation\n(%)",
        "I": "Farm Area\nHigh MAP\n\n(km²)",
        "J": "Farm Area\nLow MAP\n\n(km²)",
        "K": "Hi/Lo flow ratio fragmen\ntation (%)",
        "L": "External Fragmen\ntation copied (%)",
        "M": "Selected Fragmen\ntation\n(%)",
        "N": "Upstream\nInflow\nAbove Dam\n(%)",
        "O": "Farm runoff\nAbove\n (into) Dam\n(%)",
        "P": "Composite\nDam\nCapacity\n(m³)",
        "Q": "Initial Dam Volume (%)",
        "R": "Min. Dam Volume for Transfers (%)",
        "S": "Downstream diversion \nback to dam\n(m³/s)",
        "T": "Downstream diversion \nback to dam\n(m³/day)",
        "U": "Irrigation\nReturn\nFlow \n(%)",
    }
    for c, h in headers.items():
        put(ws, f"{c}30", h)
    total_area = sum(f["total"] for f in FARM_SPEC)
    sum_hi = sum(f["hi"] for f in FARM_SPEC)
    sum_lo = sum(f["lo"] for f in FARM_SPEC)
    r = 31
    for i, f in enumerate(FARM_SPEC):
        put(ws, f"D{r}", f["name"])
        put(ws, f"E{r}", "Farm")
        put(ws, f"F{r}", i + 1)
        put(ws, f"G{r}", f["total"])
        put(ws, f"H{r}", f["total"] / total_area)
        put(ws, f"I{r}", f["hi"])
        put(ws, f"J{r}", f["lo"])
        put(ws, f"K{r}", f["hi"] / sum_hi * HI_LO_SPLIT[0] + f["lo"] / sum_lo * HI_LO_SPLIT[1])
        for c, k in (("L", "ext"), ("M", "ext"), ("N", "up"), ("O", "runoff"), ("P", "cap"), ("Q", "init"), ("R", "min"), ("T", "divert"), ("U", "ret")):
            if f[k] is not None:
                put(ws, f"{c}{r}", f[k])
        divert = float(f["divert"])
        if divert:
            put(ws, f"S{r}", divert / 86400)
        r += 1
    sentinel = r
    put(ws, f"D{sentinel}", "-- do NOT delete this row !  Insert here")
    for c in "EFGHIJKLMNOPQRSTU":
        put(ws, f"{c}{sentinel}", "--")
    put(ws, f"G{sentinel + 1}", total_area)
    put(ws, f"I{sentinel + 1}", sum_hi)
    put(ws, f"J{sentinel + 1}", sum_lo)
    put(ws, f"M{sentinel + 1}", sum(f["ext"] or 0 for f in FARM_SPEC))
    define(wb, "zFarmSpec_FarmNameLst", s, ref("D", 30, "D", sentinel))
    define(wb, "rFarmSpec_AreaTotal", s, ref("G", sentinel + 1))
    define(wb, "rFarmSpec_DataAreas", s, ref("I", 30, "J", sentinel))
    define(wb, "rFarmSpec_DataDam", s, ref("O", 30, "Q", sentinel))
    define(wb, "rFarmSpec_FragmentationTolerance", s, ref("C", 16))
    define(wb, "rFarmSpec_FragmentationTotal", s, ref("M", sentinel + 1))
    define(wb, "rFarmSpec_Methods", s, ref("M", 4, "M", 6))
    define(wb, "rFarmSpec_MethodArea", s, ref("M", 4))
    define(wb, "rFarmSpec_MethodHiLo", s, ref("M", 5))
    define(wb, "rFarmSpec_SelectedMethod", s, ref("M", 27))
    for name, c in (
        ("zFarmSpec_PercFragmLst", "M"),
        ("zFarmSpec_PercUpstrInflowToDamLst", "N"),
        ("zFarmSpec_PercFarmRunoffToDamList", "O"),
        ("zFarmSpec_CompositeDamVol", "P"),
        ("zFarmSpec_StartStoragePercLst", "Q"),
        ("zFarmSpec_CompositeDamMinPerc", "R"),
        ("zFarmSpec_DiversionToDam", "T"),
        ("zFarmSpec_PercIrrReturnFlow", "U"),
    ):
        define(wb, name, s, ref(c, 30, c, sentinel))


def crop_demand(wb) -> None:
    s = "Crop demand"
    ws = wb.create_sheet(s)
    put(ws, "A1", "Crop demand")
    put(ws, "D12", "Calculation precision: ")
    put(ws, "F12", 2)
    put(ws, "D14", "Month: ")
    for i, m in enumerate(MONTH_NUMS_WY):
        put(ws, f"{col('F', i)}14", m)
    put(ws, "D19", "Rainfall & evaporation (mm)")
    for i, m in enumerate(MONTHS_WY):
        put(ws, f"{col('F', i)}19", m)
    put(ws, "D20", "Synthetic A-pan Evap")
    for i, v in enumerate(APAN_MM):
        put(ws, f"{col('F', i)}20", v)
    put(ws, "D21", "-- do NOT delete this row !  Insert above")
    put(ws, "D23", "Effective rainfall % : ")
    put(ws, "F23", EFFECTIVE_RAIN)
    put(ws, "E28", "Factors")
    put(ws, "R28", "Gross")
    put(ws, "F29", "Crop Factors")
    put(ws, "S29", "Crop irrigation demand, Gross (mm)")
    put(ws, "D30", "Crop name definitions")
    for i, m in enumerate(MONTHS_WY):
        put(ws, f"{col('F', i)}30", m)
        put(ws, f"{col('S', i)}30", m)
    r = 31
    for name, factors in CROPS:
        put(ws, f"D{r}", CROP_AS_TYPED.get(name, name))
        for i, f in enumerate(factors):
            put(ws, f"{col('F', i)}{r}", f)
            put(ws, f"{col('S', i)}{r}", round(APAN_MM[i] * f, 2))
        r += 1
    sentinel = r
    put(ws, f"D{sentinel}", "-- do NOT delete this row !  Insert here")
    for i in range(12):
        put(ws, f"{col('F', i)}{sentinel}", "--")
        put(ws, f"{col('S', i)}{sentinel}", "--")
    define(wb, "zCropDemand_CropNameLst", s, ref("D", 30, "D", sentinel))
    define(wb, "zCropDemand_FactorsTbl", s, ref("E", 30, "R", sentinel))
    define(wb, "zCropDemand_GrossTbl", s, ref("R", 30, "AE", sentinel))
    define(wb, "rCropDemand_EffectiveRainfall", s, ref("F", 23))
    define(wb, "rCropDemand_CalcPrecision", s, ref("F", 12))


def farm_demand(wb) -> None:
    s = "Farm demand"
    ws = wb.create_sheet(s)
    put(ws, "A1", "Farm demand")
    put(ws, "D19", "Calculation")
    put(ws, "E19", "precision: ")
    put(ws, "F19", 1)
    put(ws, "P17", "Month# ")
    put(ws, "P19", "Days/Month ")
    for i, (m, d) in enumerate(zip(MONTH_NUMS_WY, DAYS_WY)):
        put(ws, f"{col('Q', i)}17", m)
        put(ws, f"{col('Q', i)}19", d)
    put(ws, "G23", "Farm Crop area (m²)")
    put(ws, "Q23", "Farm gross irrigation demand (m³ / day)")
    put(ws, "D24", "Farm name")
    put(ws, "E24", "Type check")
    put(ws, "F24", "Sequence check")
    put(ws, "G24", "|")
    last_crop = col("H", len(FARM_DEMAND_CROPS) - 1)
    sep = col(last_crop, 1)
    total_col = col(sep, 1)
    for i, c in enumerate(FARM_DEMAND_CROPS):
        put(ws, f"{col('H', i)}24", c)
    put(ws, f"{sep}24", "|")
    put(ws, f"{total_col}24", "Total")
    for i, m in enumerate(MONTHS_WY):
        put(ws, f"{col('Q', i)}24", m)
    factors = dict(CROPS)
    r = 25
    seq = 0
    for row in FARM_DEMAND:
        if row is None:
            r += 1
            continue
        name, areas = row
        seq += 1
        put(ws, f"D{r}", name)
        put(ws, f"E{r}", "Farm")
        put(ws, f"F{r}", seq)
        put(ws, f"G{r}", "|")
        for i, a in enumerate(areas):
            if a is not None:
                put(ws, f"{col('H', i)}{r}", a)
        put(ws, f"{sep}{r}", "|")
        put(ws, f"{total_col}{r}", sum(a or 0 for a in areas))
        for m in range(12):
            gross = sum(
                (a or 0) * factors[c][m] * APAN_MM[m] / 1000 / DAYS_WY[m]
                for c, a in zip(FARM_DEMAND_CROPS, areas)
                if c in factors
            )
            put(ws, f"{col('Q', m)}{r}", TYPED_GROSS_M3_DAY.get(name, round(gross, 1)))
        r += 1
    sentinel = r
    put(ws, f"D{sentinel}", "-- do NOT delete this row !  Insert here")
    for c in ["E", "F", "G", *[col("H", i) for i in range(len(FARM_DEMAND_CROPS))], sep, total_col]:
        put(ws, f"{c}{sentinel}", "--")
    for i in range(12):
        put(ws, f"{col('Q', i)}{sentinel}", "--")
    define(wb, "zFarmDemand_FarmNameLst", s, ref("D", 24, "D", sentinel))
    define(wb, "zFarmDemand_CropNameLst", s, ref("G", 24, sep, 24))
    define(wb, "zFarmDemand_FarmCropAreaTotalLst", s, ref(total_col, 24, total_col, sentinel))
    define(wb, "zFarmDemand_GrossMth", s, ref("Q", 24, "AB", sentinel))
    define(wb, "zFarmDemand_GrossMthDays", s, ref("Q", 19, "AB", 19))


def transfers(wb) -> None:
    s = "Transfers"
    ws = wb.create_sheet(s)
    put(ws, "A1", "Transfers")
    put(ws, "I1", "    This sheet is uniquely configured for this catchment, by the Hydrologist user")
    in_out = [f for f, _ in IN_OUT]
    # I '|', one InOut column per farm, '|' (also the From block's label column), one column per transfer, '|'.
    io_first = "J"
    io_sep = col(io_first, len(in_out))
    tr_first = col(io_sep, 1)
    tr_cols = [col(tr_first, i) for i in range(len(TRANSFERS))]
    tr_end = col(tr_first, len(TRANSFERS))
    letters = {f"T{i}": c for i, c in enumerate(tr_cols)}

    put(ws, "A7", "Formula: ")
    put(ws, "H7", 1)
    put(ws, "A8", "**")
    put(ws, "B8", "As a precaution, formula are copied as text here, in case Elm. sheets are deleted!")
    for c in ("I", io_sep, tr_end):
        put(ws, f"{c}7", "|")
        put(ws, f"{c}8", "|")
    for i, (farm, template) in enumerate(IN_OUT):
        formula = template.format(**letters)
        c = col(io_first, i)
        put(ws, f"{c}7", formula)
        put_text(ws, f"{c}8", formula)
    for i, (c, (frm, *_rest)) in enumerate(zip(tr_cols, TRANSFERS)):
        formula = "=0" if i in TRANSFERS_OFF else f"=IF(fIsMthIn($H7, {c}$14), fGetTrfVolCapped('{frm}'!$Q6,{c}$10,{c}$11,{c}$16), 0)"
        put(ws, f"{c}7", formula)
        put_text(ws, f"{c}8", formula)

    labels = {
        10: "Dam capacity (m3): ",
        11: "Min capacity (%): ",
        12: "Transfer To: ",
        13: "Direction: ",
        14: "Transfer months: ",
        15: "Max Transfer to Farm (m³/s): ",
        16: "Transfer capacity max: ",
    }
    for r, lab in labels.items():
        put(ws, f"{io_sep}{r}", lab)
    caps = {f["name"]: f["cap"] or 0 for f in FARM_SPEC}
    for c, (frm, to, months, rate, min_pct, daily) in zip(tr_cols, TRANSFERS):
        put(ws, f"{c}10", caps.get(frm, 0))
        put(ws, f"{c}11", min_pct)
        put(ws, f"{c}12", to)
        put(ws, f"{c}13", f"{frm[:3]}→{to[:3]}")
        put(ws, f"{c}14", months)
        put(ws, f"{c}15", rate)
        put(ws, f"{c}16", daily)

    put(ws, "I18", "Transfer flow InOut   (m³/d);  red flag left")
    put(ws, f"{io_sep}18", "Transfers Draw From Dam  (m³/d, day -1) ;  red flag left")
    put(ws, "G19", "Date")
    put(ws, "H19", "Month")
    for c in ("I", io_sep, tr_end):
        put(ws, f"{c}19", "|")
        put(ws, f"{c}20", "|")
    for i, farm in enumerate(in_out):
        put(ws, f"{col(io_first, i)}19", farm)
    for c, (frm, *_rest) in zip(tr_cols, TRANSFERS):
        put(ws, f"{c}19", frm)
    put(ws, "G20", "init")
    define(wb, "zTransfers_HeaderDate", s, ref("G", 19))
    define(wb, "zTransfers_HeaderFarmsInOut", s, ref("I", 19, io_sep, 19))
    define(wb, "zTransfers_HeaderFarmsFrom", s, ref(io_sep, 19, tr_end, 19))
    define(wb, "zTransfers_HeaderTransfersTbl", s, ref("G", 19, tr_end, 19))
    define(wb, "zTransfers_FormulaRow", s, ref("H", 7, tr_end, 7))
    define(wb, "zTransfers_FormulasAsTxt", s, ref("I", 8, tr_end, 8))


# --------------------------------------------------------------------------
# Daily series.


def daily_series() -> dict[str, list[Any]]:
    """Deterministic invented daily values, one list per [Flow data] column."""
    rng = random.Random(SEED)
    out: dict[str, list[Any]] = {k: [] for k in ("pitman", "gauge", "logger", "rain", "chirps", "forecast")}
    q = 0.4  # m3/s
    for i, d in enumerate(days()):
        wet = d.month in WET_MONTHS
        rain = round(rng.expovariate(1 / (14.0 if wet else 6.0)), 1) if rng.random() < (0.35 if wet else 0.06) else 0.0
        chirps = round(max(0.0, rain * rng.uniform(0.6, 1.3) + (rng.uniform(0, 3) if rng.random() < 0.1 else 0.0)), 1)
        q = q * 0.96 + 0.004 * rain + 0.002
        gauge_true = round(q, 3)
        noise = rng.uniform(0.95, 1.1)

        if within(d, ZERO_RUN_WET) or within(d, ZERO_RUN_DRY):
            r_cell: Any = 0.0
        elif within(d, RAIN_BLANK):
            r_cell = None
        elif d == RAIN_TEXT_DAY:
            r_cell = "err"
        else:
            r_cell = rain
        out["rain"].append(r_cell)
        out["chirps"].append(chirps)

        if within(d, GAUGE_BLANK):
            g_cell: Any = None
        elif d == GAUGE_TEXT_DAY:
            g_cell = "n/a"
        elif d >= GAUGE_SCALED_FROM:
            g_cell = gauge_true * GAUGE_SCALE  # a scaled record: no longer whole thousandths
        else:
            g_cell = gauge_true
        out["gauge"].append(g_cell)
        out["logger"].append(round(gauge_true * noise, 3) if d >= LOGGER_FROM else None)
        out["pitman"].append(round(gauge_true * 1.2, 3) if i < PITMAN_DAYS else None)
        out["forecast"].append(None)  # an empty column is not imported
    return out


def flow_data(wb, series: dict[str, list[Any]]) -> None:
    s = "Flow data"
    ws = wb.create_sheet(s)
    ds = days()
    last = END
    put(ws, "A1", "Flow data")
    put(ws, "I2", OUTFLOW_GAUGE)
    put(ws, "B7", "First date in Date-series")
    put(ws, "C7", dt.datetime.combine(START, dt.time()))
    put(ws, "B12", "Last date in Date-series")
    put(ws, "C12", dt.datetime.combine(last, dt.time()))
    put(ws, "N7", "Convert Sec To Day : ")
    put(ws, "P7", 86400)
    put(ws, "P13", USE_FLOW)
    put(ws, "F16", "Flow A")
    put(ws, "G16", "Flow B")
    put(ws, "H16", "Flow C")
    put(ws, "O16", "Hydrologist set to use: ")
    put(ws, "P16", {1: "Flow A", 2: "Flow B", 3: "Flow C"}[USE_FLOW])
    headers = {
        "E": "Date",
        "F": "Pitman \n(Natural) \nFlow \n(m³/s)",
        "G": "Flow \nGauge Z1 \n(m³/s)",
        "H": "Logger \nL1 \n(Curr. Day) \n(m³/s)",
        "I": "Average Catchment Rainfall \n(mm)",
        # J (CHIRPS) left without a header: the series is named after its kind.
        "K": "Forecast\nRainfall\n(mm)",
        "L": "|",
        "M": "Mth#",
        "N": "Is Summer?",
        "O": "Pitman\nDaily Flow\n(Natural)\n(m³/day)",
        "P": "Observed\nDaily Flow\n(Curr. Day)\n(m³/day)",
        "R": "Use rain",
        "S": "Underlying Base Flow",
        "AB": "Resultant\n(Natural)\nflow\n(m³/day)",
        "AF": "Pragmatic EWR\n(m³/day)",
        "AG": f"Simulated outflow\n({OUTFLOW_GAUGE})\n(m³/day)",
        "AK": "EWR volume of NOT met (m³/day)",
        "AN": "|",
    }
    for c, h in headers.items():
        put(ws, f"{c}19", h)
    put(ws, "E20", "reserved")
    put(ws, "S20", BASE_FLOW_INITIAL)
    summer = {10, 11, 12, 3}
    for i, d in enumerate(ds):
        r = 21 + i
        put(ws, f"E{r}", dt.datetime.combine(d, dt.time()))
        for c, k in (("F", "pitman"), ("G", "gauge"), ("H", "logger"), ("I", "rain"), ("J", "chirps"), ("K", "forecast")):
            v = series[k][i]
            if v is not None:
                put(ws, f"{c}{r}", v)
        put(ws, f"L{r}", "|")
        put(ws, f"M{r}", d.month)
        put(ws, f"N{r}", 1 if d.month in summer else 0)
        rain = series["rain"][i]
        put(ws, f"R{r}", rain if isinstance(rain, float) else 0)
        g = series["gauge"][i]
        if isinstance(g, float):
            put(ws, f"P{r}", round(g * 86400))
        put(ws, f"AF{r}", EWR_M3_DAY[MONTH_NUMS_WY.index(d.month)])
    define(wb, "zFlowData_HeaderDate", s, ref("E", 19))
    define(wb, "zFlowData_HeaderHydrologyData", s, ref("E", 19, "L", 19))
    define(wb, "zFlowData_HeaderTbl", s, ref("E", 19, "AN", 19))
    define(wb, "zFlowData_Date1_DateSeries", s, ref("C", 7))
    define(wb, "zFlowData_DateE_DateSeries", s, ref("C", 12))
    define(wb, "zFlowData_HeaderResultFlow", s, ref("AB", 19))
    define(wb, "zFlowData_HeaderEWRpragma", s, ref("AF", 19))
    define(wb, "zFlowData_HeaderSimulated", s, ref("AG", 19))
    define(wb, "zFlowData_HeaderUseRain", s, ref("R", 19))
    define(wb, "zFlowData_HeaderObserved", s, ref("P", 19))
    define(wb, "zFlowData_HeaderEWRvolumeNotMet", s, ref("AK", 19))
    define(wb, "rUseFlow", s, ref("P", 13))


def flow_calibration(wb) -> None:
    s = "Flow Calibration Cfg"
    ws = wb.create_sheet(s)
    put(ws, "A1", "Flow Calibration Cfg")
    labels = {
        "C7": "Calibration Month Start:",
        "C8": "Calibration Grap Year1 : ",
        "C9": "Calibration YearN : ",
        "C11": "Calibration Date1 : ",
        "C12": "Calibration DateN : ",
        "G8": "Convert rain to flow if : > ",
        "G9": "a = * ",
        "G10": "b = rain ^ ",
        "G11": "Summer months: ",
        "G12": "Summer response factor = * ",
        "G13": "Winter response factor = * ",
        "G16": "Rain event: Shift Peak Index Min ",
        "G17": "Shift Peak Index Max ",
        "G18": "Min ratio Rain/Base to reset Base flow ",
        "G19": "Winter starts day+1 if rain >= ",
        "G20": "Winter starts day+0 if rain >= ",
        "G22": "Gauge Element to calibrate to: ",
        "AC4": "Set (define) Base for Peakflow: ",
        "AC5": "Rain ð Flow amplitude: ",
        "AC6": "Max. Days in curve data set: ",
        "AC7": "Lowest allowable index: ",
        "AF12": "Smoothed",
        "AF13": "Recession",
        "AG12": "Smoothed",
        "AG13": "flow",
        "AH12": "Day",
        "AH13": "#",
    }
    for c, v in labels.items():
        put(ws, c, v)
    for c, v in CALIBRATION.items():
        put(ws, c, v)
    for k in range(80):
        r = 14 + k
        put(ws, f"AF{r}", round(0.3 + 0.68 * (1 - math.exp(-k / 6)), 4))
        put(ws, f"AG{r}", 0 if k == 79 else round(900000 * math.exp(-k / 8)))
        put(ws, f"AH{r}", k + 1)
    for name, cell in (
        ("rCalibration_Amplitude", "AD5"),
        ("rCalibration_BasePeak", "AD4"),
        ("rCalibration_DaysMax", "AD6"),
        ("rCalibration_MinIndex", "AD7"),
        ("rCalibration_Mth1", "D7"),
        ("rCalibration_FactorSummer", "H12"),
        ("rCalibration_FactorWinter", "H13"),
        ("rCalibration_MinFlowRatioToResetBase", "H18"),
        ("rCalibration_PeakCoef_a", "H9"),
        ("rCalibration_PeakExp_b", "H10"),
        ("rCalibration_RainThreshold", "H8"),
        ("rCalibration_ShiftPeakIndexHi", "H17"),
        ("rCalibration_ShiftPeakIndexLo", "H16"),
        ("rCalibration_SummerMths", "H11"),
        ("rCalibration_WinterNextDayThresh", "H19"),
        ("rCalibration_WinterTodayThresh", "H20"),
        ("zCalibration_Date1", "D11"),
        ("zCalibration_DateN", "D12"),
        ("zCalibration_GaugeToCalibrate", "H22"),
    ):
        m = re.match(r"([A-Z]+)(\d+)", cell)
        define(wb, name, s, ref(m[1], int(m[2])))
    define(wb, "rCalibration_RecessionFactors", s, ref("AF", 14, "AF", 93))
    define(wb, "rCalibration_Curve", s, ref("AG", 14, "AG", 93))


def ewr_cfg(wb) -> None:
    s = "EWR Cfg"
    ws = wb.create_sheet(s)
    put(ws, "F80", "Pragmatic EWR")
    put(ws, "G87", "Month#")
    put(ws, "J87", "Pragmatic EWR flow m³/day")
    put(ws, "K87", "m³/s")
    for i, (lab, m, v) in enumerate(zip(MONTHS_WY, MONTH_NUMS_WY, EWR_M3_DAY)):
        r = 88 + i
        put(ws, f"F{r}", lab)
        put(ws, f"G{r}", m)
        put(ws, f"J{r}", v)
        put(ws, f"K{r}", v / 86400)
    define(wb, "zEWR_Pragmatic", s, ref("J", 88, "J", 99))


def ewr_pivot(wb) -> None:
    s = "EWR shortfalls Pivot Data"
    ws = wb.create_sheet(s)
    for c, h in zip("DEFGHI", ["Year", "Month", "Farm", "Vol Month", "Cnt Not Met / \nMth   .", "% Not Met /\nMth   ."]):
        put(ws, f"{c}19", h)
    rows = [(2019, 10, f, 0, 0) for f in farm_names()] + [(2019, 11, "Golf Farm", -1250, 4), (2019, 11, "Hotel Farm", -310, 1)]
    for i, (y, m, f, vol, cnt) in enumerate(rows):
        r = 20 + i
        for c, v in zip("DEFGH", (y, m, f, vol, cnt)):
            put(ws, f"{c}{r}", v)
        put(ws, f"I{r}", cnt / 31)
    define(wb, "zEWRshortPivotData_DataTbl", s, ref("D", 19, "I", 19 + len(rows)))


def shortfalls(wb) -> None:
    s = "Shortfalls"
    ws = wb.create_sheet(s)
    put(ws, "G15", "Set reporting period: ")
    put(ws, "H15", dt.datetime.combine(START, dt.time()))
    put(ws, "H16", dt.datetime.combine(END, dt.time()))
    headers = {"E": "Column\n in EWR short\nfalls", "G": "Farm name", "H": "Average Demand\n\n(m³/day)", "I": "Average Supply\n\n(m³/day)",
               "J": "Average deficit\n\n(m³/day)", "K": "%Supply\n\n\n(%)", "M": "Target Volume\n\n(m³/day)", "N": "- = Reduce\n+ = Gain\n\n(m³/day)",
               "O": "- = Reduce\n+ = Gain\n\n(l/sec)", "P": "%Supply\n\n\n(%)", "R": "Average \nEWR shortfall \nfor the period\n(m³/day)",
               "S": "Total reductions\n\n(m³/day)", "T": "Total reductions\n\n(l/sec)",
               "U": "Volume left after irrigation balanced & EWR met   (m³/day)", "V": "Reduction of demand required\n(%)"}
    for c, h in headers.items():
        put(ws, f"{c}20", h)
    rng = random.Random(SEED + 1)
    farms = farm_names()
    tot = {c: 0.0 for c in "HIJMNORSU"}
    for i, f in enumerate(farms):
        r = 21 + i
        demand = 0 if f in ("Charlie Farm", "Delta Farm") else round(rng.uniform(400, 4000))
        supplied = round(demand * rng.uniform(0.7, 1.0))
        ewr = -round(rng.uniform(0, 300)) if i % 3 == 0 else 0
        vals = {"H": demand, "I": supplied, "J": supplied - demand, "M": supplied, "N": 0, "O": 0, "R": ewr, "S": ewr, "T": round(ewr / 86.4, 1), "U": supplied + ewr}
        put(ws, f"E{r}", get_column_letter(12 + i))
        put(ws, f"G{r}", f)
        for c, v in vals.items():
            put(ws, f"{c}{r}", v)
            if c in tot:
                tot[c] += v
        for c in ("K", "P", "V"):
            put(ws, f"{c}{r}", round(supplied / demand, 3) if demand else "-")
    sentinel = 21 + len(farms)
    put(ws, f"G{sentinel}", "-- do NOT delete this row !  Insert here")
    for c in "EHIJKMNOPRSTUV":
        put(ws, f"{c}{sentinel}", "--")
    for c, v in tot.items():
        put(ws, f"{c}{sentinel + 1}", v)
    put(ws, f"K{sentinel + 1}", round(tot["I"] / tot["H"], 3))
    define(wb, "zShortfalls_Tbl", s, ref("E", 20, "V", sentinel))
    define(wb, "zShortfalls_PeriodStart", s, ref("H", 15))


FARM_HEADERS = {
    "F": "Irrigation \nDemand\nNet\n(m³/d)     .",
    "G": "Irrigation\nActual\n\n(m³/d)     .",
    "H": "Upstream\nInflow\n\n(m³/d)     .",
    "I": "Farm\nRunoff\n\n(m³/d)     .",
    "J": "Transfer\nin  +ve\nout  -ve\n(m³/d)     .",
    "Q": "Storage\nEnd of day\n\n(m³/d)    .",
    "R": "Spill\n\n\n(m³/d)    .",
    "U": "Farm\nOutflow\n\n(m³/d)     .",
    "W": "Irrigation \nDeficit\n\n(m³/d)     .",
    "Y": "Fragmented\nEWR\n\n(m³/d)     .",
    "Z": "Cummulative\nEWR\n\n(m³/d)     .",
    "AA": "Cummulative\nEWR\nshortfall (-)\n(m³/d)     .",
    "AB": "Incremental \nEWR \nshortfall (-)\n(m³/d)     .",
}
GAUGE_HEADERS = {
    "F": "Copy in measured Flow data\n(m³/d)     .",
    "G": "Upstream\nInflow\n(=outflow)\n(m³/d)     .",
    "H": "Cummulative\nEWR\n\n(m³/d)     .",
    "I": "Upstr Cumm\nThis G. Incr\nEWR short\n(m³/d)     .",
}
ELEMENT_DAYS = 7  # Element sheets hold a week of results: only expected.json reads them
ELEMENT_START_OFFSET = 0  # days after [Flow data]'s first date that the Element sheets (the model window) start


def element_sheets(wb) -> None:
    rng = random.Random(SEED + 2)
    for name, etype, _ in NETWORK:
        clean = re.sub(r"\s+", " ", name)
        if clean in NO_ELEMENT_SHEET:
            continue
        ws = wb.create_sheet(clean)
        headers = FARM_HEADERS if etype == "Farm" else GAUGE_HEADERS
        put(ws, "B15", "Reserved")
        put(ws, "E19", "Date")
        for c, h in headers.items():
            put(ws, f"{c}19", h)
        put(ws, "E20", "init")
        for i in range(ELEMENT_DAYS):
            r = 21 + i
            put(ws, f"E{r}", dt.datetime.combine(START + dt.timedelta(days=ELEMENT_START_OFFSET + i), dt.time()))
            for c in headers:
                put(ws, f"{c}{r}", round(rng.uniform(-500, 5000)))


# --------------------------------------------------------------------------


def build() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    app_settings(wb)
    home(wb)
    templates(wb)
    network(wb)
    farm_spec(wb)
    crop_demand(wb)
    farm_demand(wb)
    transfers(wb)
    flow_data(wb, daily_series())
    flow_calibration(wb)
    ewr_cfg(wb)
    ewr_pivot(wb)
    shortfalls(wb)
    element_sheets(wb)
    wb.properties.creator = "scripts/wbt-import/make_synthetic_workbook.py"
    wb.properties.lastModifiedBy = None
    wb.properties.title = "Synthetic b023 workbook (invented catchment)"
    wb.properties.created = FIXED_TIME
    return wb


def write_workbook(path: Path) -> None:
    """Save the workbook with fixed timestamps (core.xml and zip entries), so the bytes are reproducible."""
    buf = io.BytesIO()
    build().save(buf)
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ")
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename != "[Content_Types].xml"):
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + stamp.encode() + rb"\g<2>", data)
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, data)
    path.write_bytes(out.getvalue())


def dump_project(project: dict[str, Any]) -> str:
    return json.dumps(project, indent=2, ensure_ascii=False) + "\n"


def dump_notes(notes: list[str]) -> str:
    return "".join(f"{n}\n" for n in notes)


def expected_outputs(workbook: Path) -> dict[str, str]:
    """The fixture text files for a workbook: {file name: content}."""
    out = {}
    for suffix, kwargs in (("", {}), (".gauge-reference", GAUGE_REFERENCE_ARGS), (".run-of-river", {"run_of_river": True})):
        project, _expected, notes = extract(workbook, **kwargs)
        out[f"{STEM}{suffix}.project.json"] = dump_project(project)
        out[f"{STEM}{suffix}.notes.txt"] = dump_notes(notes)
    return out


def main(argv: list[str]) -> int:
    outdir = Path(argv[1]) if len(argv) > 1 else FIXTURES
    outdir.mkdir(parents=True, exist_ok=True)
    workbook = outdir / WORKBOOK_NAME
    write_workbook(workbook)
    files = expected_outputs(workbook)
    for name, text in files.items():
        (outdir / name).write_text(text, encoding="utf-8")
    print(f"{workbook} ({workbook.stat().st_size} bytes) + {', '.join(files)}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
