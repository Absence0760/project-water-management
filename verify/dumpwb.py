import openpyxl,sys
from openpyxl.worksheet.formula import ArrayFormula
wb=openpyxl.load_workbook(sys.argv[1],data_only=False,read_only=True,keep_vba=False)
maxrows=int(sys.argv[2]) if len(sys.argv)>2 else 40
only=sys.argv[3].split(',') if len(sys.argv)>3 else None
for ws in wb.worksheets:
    if only and ws.title not in only: continue
    print('#####',ws.title, ws.max_row, ws.max_column)
    for r,row in enumerate(ws.iter_rows(max_row=maxrows),1):
        cells=[]
        for c in row:
            v=c.value
            if v is None or not hasattr(c,'coordinate'): continue
            if isinstance(v,ArrayFormula): v='{ARR '+str(v.ref)+'}'+str(v.text)
            cells.append(f'{c.coordinate}={v!r}')
        if cells: print(' | '.join(cells))
